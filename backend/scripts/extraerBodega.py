"""
Extrae a JSON los Excel de la bodega del centro (plantillas/).

    python scripts/extraerBodega.py <carpeta plantillas> <salida.json>

Solo lee y transcribe: la limpieza y el emparejamiento de nombres los hace
scripts/importarBodega.js con funciones probadas. Se leen los valores ya
calculados (data_only), porque los nombres y saldos de muchas hojas son
fórmulas que apuntan a otros libros.
"""
import json
import re
import sys
import datetime
from pathlib import Path

import openpyxl

MESES = {'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6, 'JUL': 7, 'AGO': 8,
         'SEP': 9, 'SEPT': 9, 'OCT': 10, 'NOV': 11, 'DIC': 12,
         'ENERO': 1, 'FEBRERO': 2, 'MARZO': 3, 'ABRIL': 4, 'MAYO': 5, 'JUNIO': 6, 'JULIO': 7,
         'AGOSTO': 8, 'SEPTIEMBRE': 9, 'SETIEMBRE': 9, 'OCTUBRE': 10, 'NOVIEMBRE': 11, 'DICIEMBRE': 12}


def periodo_de_hoja(nombre):
    """'SEP26' -> '2026-09'; 'Enero 19' -> '2019-01'; 'Sept26' -> '2026-09'."""
    m = re.match(r'^\s*([A-Za-zÁÉÍÓÚáéíóú]+)\s*(\d{2,4})\s*$', nombre)
    if not m:
        return None
    mes = MESES.get(m.group(1).upper()) or MESES.get(m.group(1).upper()[:3])
    anio = int(m.group(2))
    if anio < 100:
        anio += 2000
    return f'{anio}-{mes:02d}' if mes else None


def valor(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def balance(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    meses = []
    for nombre in wb.sheetnames:
        periodo = periodo_de_hoja(nombre)
        if not periodo:
            continue
        ws = wb[nombre]
        filas = list(ws.iter_rows(min_row=1, max_row=330, max_col=205, values_only=True))
        servicios = [(j, str(v).strip()) for j, v in enumerate(filas[2]) if isinstance(v, str) and j >= 2]
        # Los bloques empiezan donde la fila 4 dice "Saldo anterior".
        bloques = [j for j, v in enumerate(filas[3]) if isinstance(v, str) and v.strip().lower().startswith('saldo ant')]
        totales = next((j for j, v in enumerate(filas[2]) if isinstance(v, str) and v.strip().lower() == 'totales'), None)
        nombres_bloque = []
        for j in bloques:
            if totales is not None and j >= totales:
                continue
            # El nombre del servicio esta en la fila 3, a veces corrido una o dos columnas.
            cerca = [n for (k, n) in servicios if j - 2 <= k <= j + 2]
            nombres_bloque.append((j, cerca[0] if cerca else None))
        renglones = []
        seccion = None
        for i, fila in enumerate(filas[4:], start=5):
            nombre_insumo = fila[1]
            if not isinstance(nombre_insumo, str) or not nombre_insumo.strip():
                continue
            # Un titulo de seccion ("PLANIFICACION FAMILIAR") va en mayusculas,
            # sin numero y sin datos en el primer bloque. Un insumo sin datos
            # en post-consulta no es titulo: tiene nombre en minusculas.
            if fila[0] is None and nombre_insumo.strip().isupper() and all(v is None for v in fila[2:8]):
                seccion = nombre_insumo.strip()
                continue
            r = {'fila': i, 'numero': valor(fila[0]), 'nombre': nombre_insumo.strip(), 'seccion': seccion, 'servicios': []}
            for j, nombre_servicio in nombres_bloque:
                r['servicios'].append({
                    'servicio': nombre_servicio,
                    'saldoAnterior': valor(fila[j]), 'entradas': valor(fila[j + 1]), 'entregado': valor(fila[j + 2]),
                    'noEntregado': valor(fila[j + 3]), 'saldoSiguiente': valor(fila[j + 4]), 'promedio': valor(fila[j + 5]),
                })
            if totales is not None:
                t = [valor(v) for v in fila[totales:totales + 12]]
                r['totales'] = dict(zip(['B', 'C', 'D', 'E', 'F', 'saldoPorDemanda', 'H', 'I', 'J', 'K', 'L', 'M'], t))
                r['existenciaMesAnterior'] = valor(fila[totales - 1])
            renglones.append(r)
        # El nivel maximo esta en la fila 2, a la derecha del bloque de totales.
        nivel = next((valor(v) for j, v in enumerate(filas[1]) if totales and j > totales and isinstance(v, (int, float))), None)
        meses.append({'hoja': nombre, 'periodo': periodo, 'nivelMaximo': nivel, 'renglones': renglones})
    return meses


def cierres(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    meses = []
    for nombre in wb.sheetnames:
        periodo = periodo_de_hoja(nombre)
        if not periodo:
            continue
        ws = wb[nombre]
        renglones = []
        seccion = None
        for i, fila in enumerate(ws.iter_rows(min_row=3, max_row=400, max_col=7, values_only=True), start=3):
            nombre_insumo = fila[1]
            if not isinstance(nombre_insumo, str) or not nombre_insumo.strip():
                continue
            if fila[0] is None and all(v is None for v in fila[2:6]):
                seccion = nombre_insumo.strip()
                continue
            renglones.append({
                'fila': i, 'nombre': nombre_insumo.strip(), 'seccion': seccion,
                'saldoAnterior': valor(fila[2]), 'entradas': valor(fila[3]), 'entregado': valor(fila[4]), 'saldoActual': valor(fila[5]),
            })
        meses.append({'hoja': nombre, 'periodo': periodo, 'renglones': renglones})
    return meses


def semaforo(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    hojas = {}
    for nombre in ['Rojo', 'Amarillo', 'Verde']:
        if nombre not in wb.sheetnames:
            continue
        ws = wb[nombre]
        renglones = []
        for i, fila in enumerate(ws.iter_rows(min_row=2, max_col=12, values_only=True), start=2):
            if not isinstance(fila[1], str) or not fila[1].strip():
                continue
            lotes = []
            for k in range(3):
                cantidad, lote, vence = fila[2 + k * 3], fila[3 + k * 3], fila[4 + k * 3]
                if lote is None and not cantidad:
                    continue
                lotes.append({'cantidad': valor(cantidad), 'lote': None if lote is None else str(lote).strip(), 'vence': valor(vence)})
            renglones.append({'fila': i, 'nombre': fila[1].strip(), 'lotes': lotes, 'saldo': valor(fila[11])})
        hojas[nombre] = renglones
    return hojas


def main():
    carpeta = Path(sys.argv[1])
    salida = Path(sys.argv[2])
    archivo = lambda patron: next(carpeta.glob(patron))
    datos = {
        'balance': balance(archivo('BALANCE*.xlsx')),
        'cierres': cierres(archivo('INVENTARIOS CIERRE*.xlsx')),
        'semaforo': semaforo(archivo('SEMAFORO*.xlsx')),
        'archivoSemaforo': archivo('SEMAFORO*.xlsx').name,
    }
    salida.write_text(json.dumps(datos, ensure_ascii=False), encoding='utf-8')
    print('balance', len(datos['balance']), 'meses;', 'cierres', len(datos['cierres']), 'meses;',
          'semaforo', {k: len(v) for k, v in datos['semaforo'].items()})


if __name__ == '__main__':
    main()
