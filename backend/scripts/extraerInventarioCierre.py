"""
Extrae el Excel INVENTARIOS CIERRE del centro (una hoja por mes) a JSON:
por cada hoja, el periodo y sus renglones con saldo anterior, entradas,
entregado y saldo actual. Lee los valores calculados (data_only).

    python scripts/extraerInventarioCierre.py <excel> <salida.json>
"""
import json
import re
import sys

import openpyxl

MESES = {
    'ene': 1, 'feb': 2, 'mar': 3, 'abr': 4, 'may': 5, 'jun': 6,
    'jul': 7, 'ago': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dic': 12,
}


def periodo_de(hoja):
    """'Enero 19', 'Junio26', 'Sept26' -> '2019-01', '2026-06', '2026-09'."""
    t = hoja.strip().lower()
    m = re.match(r'([a-záéíóú]+)\.?\s*(\d{2,4})$', t)
    if not m:
        return None
    mes = MESES.get(m.group(1)[:3].replace('é', 'e').replace('í', 'i'))
    if not mes:
        return None
    anio = int(m.group(2))
    if anio < 100:
        anio += 2000
    return f'{anio}-{mes:02d}'


def numero(v):
    if v is None or v == '':
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    texto = str(v).strip().replace(',', '')
    try:
        return int(round(float(texto)))
    except ValueError:
        return None


def main(ruta, salida):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    meses = []
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        if not hasattr(ws, 'iter_rows'):
            continue
        periodo = periodo_de(hoja)
        if not periodo:
            print('Hoja sin periodo reconocible:', hoja)
            continue
        renglones = []
        titulo = ''
        for fila in ws.iter_rows(values_only=True):
            celdas = list(fila[:6]) + [None] * (6 - len(fila[:6]))
            n0, nombre, ant, ent, entr, act = celdas
            if isinstance(nombre, str) and 'al ' in nombre and 'Camot' in nombre:
                titulo = nombre.strip()
            if not isinstance(nombre, str) or not nombre.strip():
                continue
            valores = [numero(x) for x in (ant, ent, entr, act)]
            if not isinstance(n0, (int, float)) or any(v is None for v in valores):
                continue
            renglones.append({
                'orden': int(n0),
                'nombre': nombre.strip(),
                'saldoAnterior': valores[0],
                'entradas': valores[1],
                'entregado': valores[2],
                'saldoActual': valores[3],
            })
        meses.append({'hoja': hoja, 'periodo': periodo, 'titulo': titulo, 'renglones': renglones})
    with open(salida, 'w', encoding='utf-8') as f:
        json.dump(meses, f, ensure_ascii=False)
    print(f'{len(meses)} hojas extraídas')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
