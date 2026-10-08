"""
Lee el balance del centro (BALANCE*.xlsx) tal como lo muestra Excel, para
compararlo con el sistema (scripts/verificarBalance.js).

    python scripts/leerBalanceExcel.py <BALANCE.xlsx> <INVENTARIOS CIERRE.xlsx> <salida.json>

Cada celda se calcula con su propia formula, como lo hace Excel al abrir el
libro: no se imita la formula "que deberia ser", porque el libro tiene
totales escritos a mano y formulas que apuntan a otra fila, y lo que el
centro ve es el resultado de esas formulas. Los vinculos a otros libros
(los balances de los puestos) toman el valor que Excel guardo.

El vinculo al libro INVENTARIOS CIERRE se rompio en algunos meses: las
entradas (C) y la existencia fisica (I) quedaron como =#REF! y todo lo que
depende de ellas sale #REF!. Esas dos celdas se toman del cierre, dos filas
mas arriba, como hace el vinculo en los meses sanos (=[2]Sept26!D3 y F3); el
resto se recalcula con sus formulas. Las celdas recuperadas quedan anotadas
en 'reconstruidos'.

Estructura de cada hoja (fila 3 nombres de servicio, fila 4 encabezados):
  - Un bloque de 6 columnas por servicio: saldo anterior, entradas,
    entregado, no entregado, saldo siguiente, promedio.
  - El bloque de totales: mes anterior, saldo anterior (B), entradas (C),
    entregado (D), no entregado (E), demanda (F), saldo x demanda,
    saldo siguiente (H), existencia fisica (I), promedio (J), meses de
    existencia (K), cantidad maxima (L), cantidad a solicitar (M).
"""
import datetime
import json
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

MESES = {'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6, 'JUL': 7, 'AGO': 8,
         'SEP': 9, 'OCT': 10, 'NOV': 11, 'DIC': 12}
NOMBRES_MES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
               'septiembre', 'octubre', 'noviembre', 'diciembre']
CAMPOS_SERVICIO = ['saldoAnterior', 'entradas', 'entregado', 'noEntregado', 'saldoSiguiente', 'promedio']
CAMPOS_TOTALES = ['mesAnterior', 'B', 'C', 'D', 'E', 'F', 'saldoPorDemanda', 'H', 'I', 'J', 'K', 'L', 'M']
MAX_FILA, MAX_COLUMNA = 400, 215
SECCIONES = {'MICRONUTRIENTES', 'MEDICAMENTOS ANTIMALÁRICOS', 'OTROS MEDICAMENTOS', 'PLANIFICACION FAMILIAR',
             'MATERIALES MEDICO QUIRÚRGICOS', 'MEDICAMENTOS PARA TUBERCULOSIS', 'MATERIALES DE LABORATORIO', 'COVID-19'}


def periodo_de_hoja(nombre):
    m = re.match(r'^\s*([A-Za-z]+)\s*(\d{2})\s*$', nombre)
    if not m or m.group(1).upper()[:3] not in MESES:
        return None
    return f'20{m.group(2)}-{MESES[m.group(1).upper()[:3]]:02d}'


def hojas_cierre(wb):
    """'2026-09' -> hoja del cierre ('Sept26', 'septiembre25', 'Enero 20')."""
    salida = {}
    for nombre in wb.sheetnames:
        m = re.match(r'^\s*([A-Za-z]+)\s*(\d{2})\s*$', nombre)
        if not m:
            continue
        mes = next((i for i, n in enumerate(NOMBRES_MES, start=1) if n.startswith(m.group(1).lower()[:4])), None)
        if mes:
            salida[f'20{m.group(2)}-{mes:02d}'] = nombre
    return salida


def es_numero(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def es_error(v):
    return isinstance(v, str) and v.startswith('#')


def limpio(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


class Libro:
    """Valores de las celdas del balance, calculados con sus formulas."""

    REF = re.compile(r"(?:\[(\d+)\])?(?:'([^']+)'|([A-Za-z][A-Za-z0-9_]*))!\$?([A-Z]{1,3})\$?(\d+)|\$?([A-Z]{1,3})\$?(\d+)")

    def __init__(self, ruta, ruta_cierre):
        valores = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(ruta, read_only=True)
        self.guardado = {}
        self.formula = {}
        for nombre in valores.sheetnames:
            if not periodo_de_hoja(nombre):
                continue
            self.guardado[nombre] = [list(f) for f in valores[nombre].iter_rows(
                min_row=1, max_row=MAX_FILA, max_col=MAX_COLUMNA, values_only=True)]
            self.formula[nombre] = [list(f) for f in formulas[nombre].iter_rows(
                min_row=1, max_row=MAX_FILA, max_col=MAX_COLUMNA, values_only=True)]
        cierre = openpyxl.load_workbook(ruta_cierre, read_only=True, data_only=True)
        self.cierre = {
            periodo: [list(f) for f in cierre[hoja].iter_rows(min_row=1, max_row=MAX_FILA, max_col=7, values_only=True)]
            for periodo, hoja in hojas_cierre(cierre).items()
        }
        self.calculado = {}
        self.reconstruidas = set()
        self.circulares = set()
        self.cache_viejo = []
        # Columnas de C e I en cada hoja: el vinculo al cierre (D y F del cierre).
        self.columnas_cierre = {}
        for hoja, filas in self.guardado.items():
            enc = [str(v).strip().lower() if isinstance(v, str) else '' for v in filas[3]]
            t0 = enc.index('mes anterior') + 1  # columna de Excel (base 1)
            self.columnas_cierre[hoja] = {t0 + 2: 3, t0 + 8: 5}

    @staticmethod
    def _celda(filas, fila, columna):
        if fila - 1 >= len(filas) or columna - 1 >= len(filas[fila - 1]):
            return None
        return filas[fila - 1][columna - 1]

    def valor(self, hoja, fila, columna):
        clave = (hoja, fila, columna)
        if clave in self.calculado:
            return self.calculado[clave]
        self.calculado[clave] = '#CIRCULAR!'
        v = self._valor(hoja, fila, columna)
        self.calculado[clave] = v
        return v

    def _valor(self, hoja, fila, columna):
        crudo = self._celda(self.formula[hoja], fila, columna)
        if not (isinstance(crudo, str) and crudo.startswith('=')):
            return crudo
        guardado = self._celda(self.guardado[hoja], fila, columna)
        if crudo.strip() == '=#REF!':
            fuente = self.columnas_cierre[hoja].get(columna)
            if fuente is None:
                return '#REF!'
            filas = self.cierre.get(periodo_de_hoja(hoja))
            v = filas[fila - 3][fuente] if filas and 0 <= fila - 3 < len(filas) else None
            self.reconstruidas.add((hoja, fila, columna))
            return v
        # Vinculo a otro libro: Excel muestra el valor que guardo.
        if re.search(r'\[\d+\]', crudo):
            return guardado
        v = self.evaluar(hoja, crudo[1:])
        # Referencia circular: Excel no la calcula y deja el ultimo valor que guardo.
        if v == '#CIRCULAR!':
            self.circulares.add((hoja, fila, columna))
            return guardado
        if es_numero(v) and es_numero(guardado) and abs(v - guardado) > 1e-6:
            # Solo pasa en la cadena de una referencia circular: Excel muestra lo guardado.
            self.cache_viejo.append((hoja, f'{get_column_letter(columna)}{fila}', guardado, v))
            return guardado
        if es_numero(v) and not es_numero(guardado):
            self.reconstruidas.add((hoja, fila, columna))
        return v

    def evaluar(self, hoja, expresion):
        # Las referencias se sustituyen por su valor y lo que queda es aritmetica.
        errores = []

        def sustituir(m):
            if m.group(4):
                otra = m.group(2) or m.group(3)
                if otra not in self.formula:
                    errores.append('#REF!')
                    return '0'
                v = self.valor(otra, int(m.group(5)), column_index_from_string(m.group(4)))
            else:
                v = self.valor(hoja, int(m.group(7)), column_index_from_string(m.group(6)))
            if es_error(v):
                errores.append(v)
                return '0'
            if v is None:
                return '0'
            if not es_numero(v):
                errores.append('#VALUE!')
                return '0'
            return f'({v!r})'

        aritmetica = self.REF.sub(sustituir, expresion)
        if errores:
            return errores[0]
        if not re.fullmatch(r'[\d\s.+\-*/()e]*', aritmetica):
            raise SystemExit(f'Formula no prevista en {hoja}: ={expresion}')
        try:
            return eval(aritmetica, {'__builtins__': {}})  # solo numeros y + - * / ( )
        except ZeroDivisionError:
            return '#DIV/0!'


def hoja(libro, nombre):
    filas = libro.guardado[nombre]
    encabezado = [str(v).strip().lower() if isinstance(v, str) else '' for v in filas[3]]
    nombres = filas[2]
    t0 = encabezado.index('mes anterior')
    bloques = []
    for j, texto in enumerate(encabezado[:t0]):
        if texto.startswith('saldo ant'):
            cerca = [str(nombres[k]).strip() for k in range(j - 2, j + 3)
                     if 0 <= k < len(nombres) and isinstance(nombres[k], str) and nombres[k].strip()]
            bloques.append((j, cerca[0] if cerca else None))
    nivel = next((v for j, v in enumerate(filas[1]) if j > t0 and es_numero(v)), None)

    def v(fila, j):
        return limpio(libro.valor(nombre, fila, j + 1))

    renglones = []
    seccion = None
    for i, f in enumerate(filas[4:], start=5):
        texto = f[1]
        if not isinstance(texto, str) or not texto.strip():
            continue
        # Los titulos de seccion. A veces tienen alguna cifra suelta en un
        # servicio, que no suma en ningun total: no son un insumo.
        if f[0] is None and texto.strip() in SECCIONES:
            seccion = texto.strip()
            continue
        renglones.append({
            'fila': i,
            'numero': limpio(f[0]),
            'nombre': re.sub(r'\s+', ' ', texto).strip(),
            'seccion': seccion,
            'servicios': [
                {'columna': j, 'servicio': s, **{c: v(i, j + k) for k, c in enumerate(CAMPOS_SERVICIO)}}
                for j, s in bloques
            ],
            'totales': {c: v(i, t0 + k) for k, c in enumerate(CAMPOS_TOTALES)},
            'reconstruidos': [c for k, c in enumerate(CAMPOS_TOTALES) if (nombre, i, t0 + k + 1) in libro.reconstruidas],
        })
    return {'nivelMaximo': nivel, 'servicios': [s for _, s in bloques], 'renglones': renglones}


def main():
    ruta, ruta_cierre, salida = sys.argv[1], sys.argv[2], sys.argv[3]
    libro = Libro(ruta, ruta_cierre)
    meses = [{'hoja': n, 'periodo': periodo_de_hoja(n), **hoja(libro, n)} for n in libro.guardado]
    meses.sort(key=lambda m: m['periodo'])
    with open(salida, 'w', encoding='utf-8') as archivo:
        json.dump(meses, archivo, ensure_ascii=False)
    errores = sum(1 for m in meses for r in m['renglones'] for x in r['totales'].values() if es_error(x))
    print(f'{len(meses)} meses: {meses[0]["periodo"]} a {meses[-1]["periodo"]}; '
          f'{sum(len(m["renglones"]) for m in meses)} renglones')
    print(f'Celdas recuperadas del vinculo roto al cierre: {len(libro.reconstruidas)}')
    print(f'Celdas de totales que siguen con error de Excel (#DIV/0!, #VALUE!): {errores}')
    print(f'Celdas con referencia circular (se usa el valor guardado): {len(libro.circulares)}')
    print(f'Celdas que dependen de la referencia circular (se usa el valor guardado): {len(libro.cache_viejo)}')
    for x in libro.cache_viejo[:15]:
        print('  ', x)


if __name__ == '__main__':
    main()
